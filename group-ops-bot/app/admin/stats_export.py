# -*- coding: utf-8 -*-
"""管理后台统计与导出（D3-C）：按群多维统计 + XLSX 下载。

硬约束：
- 纯读：只查 MCP 记录与本地名单，不写任何数据
- 列口径与机器人侧群命令导出一致（app/handlers/stats.py 同列同序），
  管理员下载与群成员收到的文件同源，便于互相校验
- 查询失败时拒绝出文件（不出空表冒充成功）
- ⭐ 2026-10-20 用户指令：**统一 XLSX，CSV 能力移除** —— 群内导出也是 xlsx
  （量子密信电脑端打不开 CSV），前后台不能一套 CSV 一套 xlsx
"""

from __future__ import annotations

import io
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from ..backend.base import Backend
from ..state import StateStore

BIZ_LABEL = {"checkin": "打卡", "summary": "总结", "report": "报数"}
STATUS_LABEL = {"NORMAL": "按时", "LATE": "迟到", "ADMIN_EDIT": "后台补录"}
SOURCE_LABEL = {"user": "群员", "admin": "管理员"}
SESSION_LABEL = {"am": "上午", "pm": "下午"}
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)

_QUERIES = {"checkin": "checkin_query_records",
            "summary": "summary_query_records",
            "report": "report_query_records"}


def resolve_period_key(state: StateStore, group_id: str, biz: str,
                       date_str: str) -> str:
    """日期（YYYY-MM-DD）-> 周期键。

    打卡 / 报数按日；总结按群的周期类型（日报 = 该日期，周报 = 所在 ISO 周）。
    日期非法 -> ValueError（调用方转 400 通俗提示）。
    """
    from ..period import parse_daily_key, period_key

    dt = parse_daily_key(date_str)
    if biz == "summary":
        row = state.get_group_config(group_id)
        period_type = "daily"
        if row is not None and "period_type" in row.keys():
            period_type = row["period_type"] or "daily"
        return period_key(period_type or "daily", dt)
    return date_str


def query_records(backend: Backend, biz: str, group_id: str,
                  period_key: str) -> list:
    """按业务查一个周期的记录（不跨业务混查）。"""
    query = getattr(backend, _QUERIES[biz])
    return query(group_id=group_id,
                 period_key=period_key).get("records") or []


def _fmt_quantity(value) -> str:
    try:
        number = float(value)
    except (TypeError, ValueError):
        return str(value)
    if number.is_integer():
        return str(int(number))
    return str(number)


def slot_labels_for(state: StateStore, group_id: str) -> dict:
    """本群时段 id -> 中文显示名（⭐ 2026-10-20 P1-2）。

    与机器人侧 stats._slot_labels 同口径：优先取群配置的时段名，
    回落产品默认（am=上午 / pm=下午），再命中不了就原样输出 session。
    ⛔ 旧实现一律「非 am 即下午」，自定义时段（slot_N「晚班」）导出全错。
    """
    labels = dict(SESSION_LABEL)
    try:
        from .board import _slots as _board_slots
        row = state.get_group_config(group_id)
        for slot in _board_slots(row) if row is not None else []:
            sid = str(slot.get("id") or "").strip()
            name = str(slot.get("name") or "").strip()
            if sid and name:
                labels[sid] = name
    except Exception:  # 配置读不到不影响统计主流程（回落默认标签）
        pass
    return labels


def records_to_rows(biz: str, records: list, period_type: str = "",
                    slot_labels: Optional[dict] = None) -> tuple:
    """记录 -> (表头, 数据行)。列口径与机器人侧导出一致。

    slot_labels：本群时段 id -> 显示名（缺省 = 产品默认 am/pm 标签）。
    """
    if biz == "checkin":
        headers = ["手机号", "姓名", "时段", "打卡时间", "状态", "录入来源",
                   "图片名"]
        labels = slot_labels or SESSION_LABEL
        rows = []
        for rec in records:
            session = rec.get("session") or ""
            rows.append([
                rec.get("phone") or "",
                rec.get("name") or "",
                labels.get(session) or SESSION_LABEL.get(session, session),
                rec.get("checkin_time") or "",
                STATUS_LABEL.get(rec.get("status"), rec.get("status") or ""),
                SOURCE_LABEL.get(rec.get("entry_source"),
                                 rec.get("entry_source") or ""),
                rec.get("image_name") or "",
            ])
        return headers, rows

    if biz == "summary":
        headers = ["手机号", "姓名", "周期类型", "内容", "录入来源", "更新时间"]
        rows = []
        for rec in records:
            rows.append([
                rec.get("phone") or "",
                rec.get("name") or "",
                rec.get("period_type") or period_type or "",
                rec.get("content") or "",
                SOURCE_LABEL.get(rec.get("entry_source"),
                                 rec.get("entry_source") or ""),
                rec.get("updated_time") or "",
            ])
        return headers, rows

    # report：动态字段宽表（每人一行，与机器人侧同口径）
    columns = []
    seen = set()
    for rec in records:
        for item in rec.get("items") or []:
            fid = item.get("field_id")
            if fid in seen:
                continue
            seen.add(fid)
            name = item.get("field_name") or ("字段%s" % fid)
            unit = item.get("unit")
            columns.append((fid, "%s（%s）" % (name, unit) if unit else name))
    columns.sort(key=lambda c: (c[0] is None, c[0]))
    headers = ["手机号", "姓名", "角色"] + [header for _, header in columns]
    rows = []
    by_phone: dict = {}
    for rec in records:
        phone = rec.get("phone")
        row = by_phone.setdefault(phone, {"name": rec.get("name") or "",
                                          "role": rec.get("role") or "",
                                          "values": {}})
        for item in rec.get("items") or []:
            row["values"][item.get("field_id")] = item.get("quantity")
    for phone in sorted(by_phone):
        row = by_phone[phone]
        rows.append([phone, row["name"], row["role"]] +
                    [_fmt_quantity(row["values"].get(fid, ""))
                     for fid, _ in columns])
    return headers, rows


def multi_sheet_xlsx(sheets: list) -> bytes:
    """多工作表工作簿（⛔ 已无 CSV 分支，全系统只出 xlsx）。

    sheets: [(工作表名, 表头, 数据行), ...]；工作表顺序即传入顺序。
    表头统一样式（深蓝底白字 + 冻结首行），与群内导出的三张表观感一致。
    """
    workbook = Workbook()
    workbook.remove(workbook.active)  # 去掉默认空白表，避免多出一张
    for name, headers, rows in sheets:
        ws = workbook.create_sheet(title=str(name)[:31])
        ws.append(list(headers))
        for row in rows:
            ws.append(list(row))
        for cell in ws[1]:
            cell.fill = _HEADER_FILL
            cell.font = _HEADER_FONT
        ws.freeze_panes = "A2"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def safe_name(text: str, limit: int = 40) -> str:
    """文件名安全化（只替换非法字符，不改中文）。"""
    return "".join(ch if ch not in '\\/:*?"<>|' else "_"
                   for ch in (text or "群"))[:limit]


def to_xlsx(biz: str, records: list, period_type: str = "",
            slot_labels: Optional[dict] = None) -> bytes:
    """XLSX 字节流。"""
    headers, rows = records_to_rows(biz, records, period_type, slot_labels)
    return multi_sheet_xlsx([(BIZ_LABEL[biz], headers, rows)])


def build_stats(state: StateStore, backend: Backend, group_id: str, biz: str,
                period_key: str, slot_labels: Optional[dict] = None) -> dict:
    """页面用的多维统计（纯读）。

    slot_labels：本群时段 id -> 显示名（缺省自动从群配置解析）。
    """
    records = query_records(backend, biz, group_id, period_key)
    try:
        members = len(list(state.list_member_tasks(group_id,
                                                   enabled_only=True)))
    except Exception:
        members = 0
    phones = {r.get("phone") for r in records if r.get("phone")}
    stats = {
        "biz": biz,
        "biz_label": BIZ_LABEL[biz],
        "period_key": period_key,
        "total": len(records),
        "people": len(phones),
        "members": members,
        "missing": max(members - len(phones), 0),
        "records": records,
        "dimensions": [],
    }
    # 时段显示名（仅打卡业务需要；summary / report 不读时段配置）
    labels = None
    if biz == "checkin":
        labels = slot_labels or slot_labels_for(state, group_id)
    headers, table_rows = records_to_rows(biz, records, "", labels)
    stats["table"] = {"headers": headers, "rows": table_rows}

    if biz == "checkin":
        status_counts = {"NORMAL": 0, "LATE": 0, "ADMIN_EDIT": 0}
        for rec in records:
            key = rec.get("status")
            if key in status_counts:
                status_counts[key] += 1
        stats["extra_label"] = "迟到人次"
        stats["extra_value"] = status_counts["LATE"]
        for key, label in (("NORMAL", "按时"), ("LATE", "迟到"),
                           ("ADMIN_EDIT", "后台补录")):
            stats["dimensions"].append(("提交状态", label, status_counts[key]))
        session_counts: dict = {}
        for rec in records:
            session = rec.get("session") or "未标注"
            session_counts[session] = session_counts.get(session, 0) + 1
        for session, count in sorted(session_counts.items(),
                                     key=lambda item: str(item[0])):
            # ⭐ 2026-10-20（P1-2）：时段维度也走本群配置名，⛔ 不写死 am/pm
            label = (session if session == "未标注" else
                     ((labels or {}).get(session)
                      or SESSION_LABEL.get(session, session)))
            stats["dimensions"].append(("打卡时段", label, count))

    elif biz == "summary":
        type_counts = {"daily": 0, "weekly": 0}
        covered = 0
        for rec in records:
            key = rec.get("period_type") or "daily"
            if key in type_counts:
                type_counts[key] += 1
            if (rec.get("updated_time") or "") != (rec.get("created_time")
                                                   or ""):
                covered += 1
        stats["extra_label"] = "覆盖提交"
        stats["extra_value"] = covered
        stats["dimensions"].append(("周期", "日报", type_counts["daily"]))
        stats["dimensions"].append(("周期", "周报", type_counts["weekly"]))
        stats["dimensions"].append(("提交方式", "其中覆盖提交", covered))

    else:  # report
        role_counts: dict = {}
        field_totals: dict = {}
        field_people: dict = {}
        for rec in records:
            role = rec.get("role") or "未填角色"
            role_counts[role] = role_counts.get(role, 0) + 1
            for item in rec.get("items") or []:
                name = item.get("field_name") or ("字段%s"
                                                  % item.get("field_id"))
                field_totals[name] = field_totals.get(name, 0) + float(
                    item.get("quantity") or 0)
                field_people.setdefault(name, set()).add(rec.get("phone"))
        total_quantity = sum(field_totals.values())
        stats["extra_label"] = "填报总量"
        stats["extra_value"] = (_fmt_quantity(total_quantity)
                                if float(total_quantity).is_integer()
                                else total_quantity)
        for role, count in sorted(role_counts.items()):
            stats["dimensions"].append(("角色", role, count))
        for name, total in sorted(field_totals.items(),
                                  key=lambda item: -item[1]):
            unit = ""
            for rec in records:
                for item in rec.get("items") or []:
                    if (item.get("field_name") or "") == name and item.get("unit"):
                        unit = item.get("unit")
                        break
                if unit:
                    break
            label = "%s（%s）" % (name, unit) if unit else name
            stats["dimensions"].append(
                ("字段", "%s，%d 人填报" % (label, len(field_people[name])),
                 _fmt_quantity(total) if float(total).is_integer()
                 else total))
    return stats


def export_filename(group_name: str, biz: str, period_key: str,
                    fmt: str) -> str:
    """下载文件名（中文，RFC 5987 编码由调用方拼接头）。"""
    return "%s%s_%s.%s" % (BIZ_LABEL[biz], safe_name(group_name), period_key,
                           fmt)


def all_export_filename(group_name: str, date_str: str, fmt: str) -> str:
    """全业务单表 / 工作簿的文件名。"""
    return "全天全部业务执行清单_%s_%s.%s" % (safe_name(group_name),
                                              date_str, fmt)
