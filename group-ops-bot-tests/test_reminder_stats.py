# -*- coding: utf-8 -*-
"""定时提醒与统计导出测试（§4.5 / §4.6，脱库替身）。

⛔ push 能力未确认（§5.2）：测试只验证「计算 + 去重 + 落库」，
发送未接通时如实记 unsent —— 不构成「提醒已送达」证据。
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from app.handlers import reminder as reminder_handler
from app.handlers import stats as stats_handler
from doubles.fakes import FakeBackend
from helpers import GROUP, OTHER_PHONE, PHONE, TZ, enable_all, make_config, make_ctx

AM = datetime(2026, 10, 13, 10, 0, tzinfo=TZ)   # 09:00 已过，14:50 未过
LATE = datetime(2026, 10, 13, 20, 0, tzinfo=TZ)  # 全部规则已过


def cfg_for(tmp_db):
    return make_config(Path(tmp_db.db_path).parent)


class ReminderBackend(FakeBackend):
    """按手机号区分完成状态的提醒后端替身。"""

    def __init__(self, completed=None, missing_users=None):
        super().__init__()
        self.completed = completed or set()   # {(phone, biz, session?)}
        self.missing_users = missing_users or set()
        self._ids = {}

    def _user_id_of(self, phone):
        if phone not in self._ids:
            self._ids[phone] = len(self._ids) + 1
        return self._ids[phone]

    def user_get(self, *, phone):
        self._record("user_get", {"phone": phone})
        if phone in self.missing_users:
            return {"user": None}
        return {"user": {"id": self._user_id_of(phone), "phone": phone,
                         "name": phone[-4:]}}

    def checkin_query_status(self, *, user_id, group_id, period_key, session):
        phone = self._phone_of(user_id)
        self._record("checkin_query_status", {"user_id": user_id,
                                              "session": session})
        done = (phone, "checkin", session) in self.completed
        return {"completed": done, "record_id": 1 if done else None,
                "checkin_time": None}

    def summary_query_status(self, *, user_id, group_id, period_key):
        phone = self._phone_of(user_id)
        done = (phone, "summary", None) in self.completed
        return {"completed": done, "record_id": 1 if done else None}

    def report_query_status(self, *, user_id, group_id, period_key):
        phone = self._phone_of(user_id)
        done = (phone, "report", None) in self.completed
        return {"completed": done, "record_id": 1 if done else None}

    def _phone_of(self, user_id):
        for phone, uid in self._ids.items():
            if uid == user_id:
                return phone
        return "unknown"


def setup_group_with_required(tmp_db):
    enable_all(tmp_db)
    tmp_db.upsert_member_task(GROUP, PHONE, role="开发")
    tmp_db.upsert_member_task(GROUP, OTHER_PHONE, role="测试")


class TestReminderCompute:
    def test_am_only_before_pm_deadline(self, tmp_db):
        setup_group_with_required(tmp_db)
        backend = ReminderBackend(completed={(PHONE, "checkin", "am")})
        ctx = make_ctx(tmp_db, backend, cfg_for(tmp_db), now=AM)
        targets = reminder_handler.compute_incomplete(
            ctx, GROUP, "checkin", "2026-10-13", sessions=["am"])
        assert [t.phone for t in targets] == [OTHER_PHONE]
        assert targets[0].session == "am"

    def test_missing_user_is_incomplete(self, tmp_db):
        setup_group_with_required(tmp_db)
        backend = ReminderBackend(completed={(PHONE, "checkin", "am")},
                                  missing_users={OTHER_PHONE})
        ctx = make_ctx(tmp_db, backend, cfg_for(tmp_db), now=AM)
        targets = reminder_handler.compute_incomplete(
            ctx, GROUP, "checkin", "2026-10-13", sessions=["am"])
        assert [t.phone for t in targets] == [OTHER_PHONE]

    def test_all_completed_no_targets(self, tmp_db):
        setup_group_with_required(tmp_db)
        backend = ReminderBackend(completed={
            (PHONE, "checkin", "am"), (OTHER_PHONE, "checkin", "am")})
        ctx = make_ctx(tmp_db, backend, cfg_for(tmp_db), now=AM)
        targets = reminder_handler.compute_incomplete(
            ctx, GROUP, "checkin", "2026-10-13", sessions=["am"])
        assert targets == []

    def test_summary_targets(self, tmp_db):
        setup_group_with_required(tmp_db)
        backend = ReminderBackend(completed={(PHONE, "summary", None)})
        ctx = make_ctx(tmp_db, backend, cfg_for(tmp_db), now=LATE)
        targets = reminder_handler.compute_incomplete(
            ctx, GROUP, "summary", "2026-10-13")
        assert [t.phone for t in targets] == [OTHER_PHONE]

    def test_format_message_has_names(self, tmp_db):
        setup_group_with_required(tmp_db)
        backend = ReminderBackend()
        ctx = make_ctx(tmp_db, backend, cfg_for(tmp_db), now=AM)
        targets = reminder_handler.compute_incomplete(
            ctx, GROUP, "checkin", "2026-10-13", sessions=["am"])
        text = reminder_handler.format_message(targets)
        assert "打卡" in text and OTHER_PHONE[-4:] in text


class TestReminderTick:
    def test_tick_records_plan_and_marks_unsent(self, tmp_db):
        setup_group_with_required(tmp_db)
        backend = ReminderBackend()
        pushes = []

        def push(group_id, text):
            pushes.append((group_id, text))
            return False, "platform push capability unconfirmed"

        summary = reminder_handler.tick(tmp_db, backend, push, LATE)
        assert summary["groups"] == 1
        assert summary["new"] > 0
        assert summary["sent"] == 0
        assert summary["unsent"] == summary["new"]
        assert pushes  # 计算确实尝试了推送
        rows = tmp_db._conn.execute(
            "SELECT * FROM reminder_plan WHERE status='unsent'").fetchall()
        assert rows and rows[0]["reason"]

    def test_tick_dedup(self, tmp_db):
        setup_group_with_required(tmp_db)
        backend = ReminderBackend()
        push = lambda g, t: (False, "unconfirmed")
        first = reminder_handler.tick(tmp_db, backend, push, LATE)
        second = reminder_handler.tick(tmp_db, backend, push, LATE)
        assert first["new"] > 0
        assert second["new"] == 0  # ⛔ 同周期同人不重复计划

    def test_no_rules_passed_no_targets(self, tmp_db):
        setup_group_with_required(tmp_db)
        backend = ReminderBackend()
        early = datetime(2026, 10, 13, 8, 0, tzinfo=TZ)  # 09:00 未过
        summary = reminder_handler.tick(
            tmp_db, backend, lambda g, t: (False, "x"), early)
        assert summary["targets"] == 0 and summary["new"] == 0


class TestExportWorkbook:
    """单业务导出构建器（第六阶段；⛔ 不再有一键导全部的组合表）。"""

    def test_checkin_sheet(self, tmp_db):
        data = stats_handler.build_export_workbook(
            "checkin", None, "2026-10-13", [{
                "record_id": 1, "phone": PHONE, "name": "甲",
                "period_key": "2026-10-13", "session": "am",
                "checkin_time": "2026-10-13 09:01:00", "image_ref": "img/1.jpg",
                "image_name": "1.jpg", "image_size": 100,
                "status": "NORMAL", "entry_source": "user"}])
        from io import BytesIO
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(data))
        assert wb.sheetnames == ["打卡"]
        header = [c.value for c in wb["打卡"][1]]
        assert header == ["手机号", "姓名", "时段", "打卡时间", "状态",
                          "录入来源", "图片名"]
        row2 = [c.value for c in wb["打卡"][2]]
        assert row2[0] == PHONE and row2[2] == "上午"
        assert row2[4] == "NORMAL" and row2[5] == "user"

    def test_summary_sheet(self, tmp_db):
        data = stats_handler.build_export_workbook(
            "summary", "daily", "2026-10-13", [{
                "record_id": 2, "phone": PHONE, "name": "甲",
                "period_key": "2026-10-13", "period_type": "daily",
                "content": "修复 3 个 bug", "created_time": "t1",
                "updated_time": "t2", "entry_source": "user"}])
        from io import BytesIO
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(data))
        assert wb.sheetnames == ["总结"]
        assert [c.value for c in wb["总结"][2]] == [
            PHONE, "甲", "daily", "修复 3 个 bug", "user", "t2"]

    def test_report_dynamic_columns_only_filled_fields(self, tmp_db):
        """动态字段：只展示有人填的字段（按 field_id 归并，快照列名）。"""
        data = stats_handler.build_export_workbook(
            "report", None, "2026-10-13", [{
                "record_id": 3, "phone": PHONE, "name": "甲",
                "period_key": "2026-10-13", "role": "开发",
                "created_time": "t3",
                "items": [{"field_id": 1, "field_name": "bug修复数",
                           "quantity": 3, "unit": None},
                          {"field_id": 2, "field_name": "代码行数",
                           "quantity": 500, "unit": "行"}]}])
        from io import BytesIO
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(data))
        assert wb.sheetnames == ["报数"]
        header = [c.value for c in wb["报数"][1]]
        # 字段 3（没人填）⛔ 不出列；单位进列名（快照）
        assert header == ["手机号", "姓名", "角色", "bug修复数",
                          "代码行数（行）"]
        assert [c.value for c in wb["报数"][2]] == [
            PHONE, "甲", "开发", 3, 500]

    def test_report_empty_records_header_only(self, tmp_db):
        data = stats_handler.build_export_workbook("report", None,
                                                   "2026-10-13", [])
        from io import BytesIO
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(data))
        assert [c.value for c in wb["报数"][1]] == ["手机号", "姓名", "角色"]
        assert wb["报数"].max_row == 1

    def test_report_snapshot_name_survives_pool_rename(self, tmp_db):
        """字段历史快照：列名取提交时快照，⛔ 不随字段池后续改名而变。"""
        data = stats_handler.build_export_workbook(
            "report", None, "2026-10-13", [{
                "phone": PHONE, "name": "甲", "role": "开发",
                "items": [{"field_id": 1, "field_name": "旧名字段",
                           "quantity": 1, "unit": None}]}])
        from io import BytesIO
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(data))
        assert [c.value for c in wb["报数"][1]][-1] == "旧名字段"
