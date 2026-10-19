# -*- coding: utf-8 -*-
"""管理后台页面测试（2026-10-19 重构：React + Vite SPA 外壳）。

架构变化（用户 2026-10-19 指令）：
- 页面由 Jinja2 服务端渲染改为「后端发外壳 + /api/admin/* 供数据」
  前端工程在 admin-ui/，构建产物落 app/admin/static/spa/
- 因此页面测试的断言对象从「渲染后的中文文案」变为「外壳契约」：
  外壳可达、会话门生效、静态产物可加载、旧路由 303 兼容
- 页面级文案不再由服务端产出，业务文案守则改为校验 API 错误提示

覆盖：
- 未登录 -> 303 /login；强制改密门；登录后四大板块外壳 200
- 旧路由兼容：/groups/{gid}/{tab} -> 303 /groups?group_id=…
  已移除的 /audit、/system -> 303 回仪表盘
- SPA 产物：index.html 引用的 js/css 全部可下载，且不依赖旧 Tabler 资产
- 仪表盘口径：聚合数值正确；BackendError -> unknown（⛔ 不当 0 不当成功）
- 成员级业务看板 /groups/{gid}/board（一天 × 全员 × 三业务）
- 移除群聊：DELETE /robot-instances/{gid} 级联删配置、幂等 404
- 统计导出：单业务 CSV/XLSX 契约不变；新增 scope=all 全天单表 + 三 sheet

⛔ 这是本地自测证据，不构成真实部署 / 平台验证（§7.2）。
"""

from __future__ import annotations

import io
import re
from datetime import datetime

import pytest
from fastapi.testclient import TestClient

from app.admin import auth as admin_auth
from app.admin.server import create_admin_app
from app.period import TZ
from doubles.fakes import FakeBackend, RecordingSender
from helpers import GROUP, PHONE, enable_all, make_config

ADMIN_USER = "admin"
ADMIN_PASS = "initial-pass-1"
NEW_PASS = "formal-pass-2"


def _bootstrap(tmp_db, monkeypatch):
    monkeypatch.setenv("GOPS_ADMIN_USERNAME", ADMIN_USER)
    monkeypatch.setenv("GOPS_ADMIN_PASSWORD", ADMIN_PASS)
    admin_auth.bootstrap_admin(tmp_db)


def make_client(tmp_db, tmp_path, backend=None):
    return TestClient(create_admin_app(
        state=tmp_db, backend=backend or FakeBackend(),
        sender=RecordingSender(), config=make_config(tmp_path)))


def login(client, password=ADMIN_PASS):
    r = client.post("/api/admin/login",
                    json={"username": ADMIN_USER, "password": password})
    assert r.status_code == 200, r.text
    return r


def logged_in_client(tmp_db, tmp_path, monkeypatch, backend=None):
    """完成 登录 -> 强制改密 -> 重登 的会话（页面门已解锁）。"""
    _bootstrap(tmp_db, monkeypatch)
    client = make_client(tmp_db, tmp_path, backend)
    login(client)
    csrf = client.get("/api/admin/me").json()["csrf_token"]
    r = client.post("/api/admin/change-password",
                    json={"old_password": ADMIN_PASS,
                          "new_password": NEW_PASS},
                    headers={"X-CSRF-Token": csrf})
    assert r.status_code == 200, r.text
    login(client, password=NEW_PASS)
    return client


def csrf_of(client):
    return client.get("/api/admin/me").json()["csrf_token"]


def assert_spa_shell(response):
    """外壳契约：挂载点 + 构建产物引用 + 禁缓存。"""
    assert response.status_code == 200
    assert '<div id="root">' in response.text
    assert '/static/spa/assets/' in response.text
    assert response.headers.get("cache-control") == "no-store"


# ---------- 页面门与会话 ----------

class TestPageGuards:
    def test_login_page_renders_shell(self, tmp_db, tmp_path, monkeypatch):
        _bootstrap(tmp_db, monkeypatch)
        client = make_client(tmp_db, tmp_path)
        r = client.get("/login")
        assert_spa_shell(r)
        # ⛔ 登录页不得内嵌任何 CSRF 令牌（无会话）
        assert "csrf_token" not in r.text

    @pytest.mark.parametrize("path", [
        "/", "/groups", "/robot-instances", "/stats", "/change-password",
        "/groups/%s/config" % GROUP,
    ])
    def test_unauthenticated_redirects_to_login(self, tmp_db, tmp_path,
                                                monkeypatch, path):
        _bootstrap(tmp_db, monkeypatch)
        client = make_client(tmp_db, tmp_path)
        # starlette 1.7 TestClient 默认跟随重定向：断言跳转必须显式关闭
        r = client.get(path, follow_redirects=False)
        assert r.status_code == 303
        assert r.headers["location"] == "/login"

    def test_must_change_gate(self, tmp_db, tmp_path, monkeypatch):
        """未改密：业务页一律 303 /change-password（⛔ 不放行）。"""
        _bootstrap(tmp_db, monkeypatch)
        client = make_client(tmp_db, tmp_path)
        login(client)  # must_change_password=True
        for path in ("/", "/groups", "/robot-instances", "/stats"):
            r = client.get(path, follow_redirects=False)
            assert r.status_code == 303, path
            assert r.headers["location"] == "/change-password", path
        assert_spa_shell(client.get("/change-password"))

    def test_change_password_page_redirects_after_done(self, tmp_db, tmp_path,
                                                       monkeypatch):
        client = logged_in_client(tmp_db, tmp_path, monkeypatch)
        r = client.get("/change-password", follow_redirects=False)
        assert r.status_code == 303
        assert r.headers["location"] == "/"

    def test_login_page_redirects_when_logged_in(self, tmp_db, tmp_path,
                                                 monkeypatch):
        client = logged_in_client(tmp_db, tmp_path, monkeypatch)
        r = client.get("/login", follow_redirects=False)
        assert r.status_code == 303
        assert r.headers["location"] == "/"

    def test_must_change_session_redirect_from_login(self, tmp_db, tmp_path,
                                                     monkeypatch):
        _bootstrap(tmp_db, monkeypatch)
        client = make_client(tmp_db, tmp_path)
        login(client)
        r = client.get("/login", follow_redirects=False)
        assert r.status_code == 303
        assert r.headers["location"] == "/change-password"


# ---------- 四大板块外壳 ----------

class TestSpaPages:
    @pytest.mark.parametrize("path", ["/", "/robot-instances", "/groups",
                                      "/stats"])
    def test_business_pages_serve_shell(self, tmp_db, tmp_path, monkeypatch,
                                        path):
        enable_all(tmp_db)
        client = logged_in_client(tmp_db, tmp_path, monkeypatch)
        assert_spa_shell(client.get(path))

    def test_shell_assets_downloadable(self, tmp_db, tmp_path, monkeypatch):
        """外壳引用的每个 js/css 都必须真的能下载（产物缺失会静默白屏）。"""
        client = make_client(tmp_db, tmp_path)
        html = client.get("/login").text
        refs = re.findall(r'(?:src|href)="(/static/spa/assets/[^"]+)"', html)
        assert refs, "外壳未引用任何构建产物"
        for ref in refs:
            r = client.get(ref)
            assert r.status_code == 200, ref
            assert len(r.content) > 1024, ref

    def test_shell_has_no_legacy_tabler_dependency(self, tmp_db, tmp_path,
                                                  monkeypatch):
        client = make_client(tmp_db, tmp_path)
        html = client.get("/login").text
        assert "tabler" not in html.lower()
        assert "jquery" not in html.lower()


# ---------- 旧路由兼容 ----------

class TestLegacyRoutes:
    @pytest.mark.parametrize("tab", ["config", "members", "whitelist",
                                     "fields"])
    def test_group_tab_redirects_to_spa(self, tmp_db, tmp_path, monkeypatch,
                                        tab):
        enable_all(tmp_db)
        client = logged_in_client(tmp_db, tmp_path, monkeypatch)
        r = client.get("/groups/%s/%s" % (GROUP, tab),
                       follow_redirects=False)
        assert r.status_code == 303
        assert r.headers["location"] == "/groups?group_id=%s" % GROUP

    def test_group_backfill_page_removed(self, tmp_db, tmp_path, monkeypatch):
        """补录已改为页面内弹窗，旧的补录页不再存在（⛔ 不做死链假成功）。"""
        enable_all(tmp_db)
        client = logged_in_client(tmp_db, tmp_path, monkeypatch)
        assert client.get("/groups/%s/backfill" % GROUP).status_code == 404

    def test_group_tab_unknown_404(self, tmp_db, tmp_path, monkeypatch):
        enable_all(tmp_db)
        client = logged_in_client(tmp_db, tmp_path, monkeypatch)
        assert client.get("/groups/%s/nope" % GROUP).status_code == 404

    def test_group_unknown_404(self, tmp_db, tmp_path, monkeypatch):
        client = logged_in_client(tmp_db, tmp_path, monkeypatch)
        assert client.get("/groups/no-such-group/config").status_code == 404

    @pytest.mark.parametrize("path", ["/audit", "/system"])
    def test_removed_pages_redirect_home(self, tmp_db, tmp_path, monkeypatch,
                                         path):
        """用户 2026-10-19 指令：审计与系统设置页不再保留（⛔ 死链）。"""
        client = logged_in_client(tmp_db, tmp_path, monkeypatch)
        r = client.get(path, follow_redirects=False)
        assert r.status_code == 303
        assert r.headers["location"] == "/"


# ---------- 仪表盘口径 ----------

class TestDashboardApi:
    def test_dashboard_aggregation(self, tmp_db, tmp_path, monkeypatch):
        enable_all(tmp_db)
        tmp_db.upsert_member_task(GROUP, PHONE, "张三", None, True)
        backend = FakeBackend()
        backend.results["checkin_query_records"] = {"records": [
            {"phone": PHONE, "name": "张三", "checkin_time":
             "2026-10-15T09:30:00+08:00", "entry_source": "user"}]}
        client = logged_in_client(tmp_db, tmp_path, monkeypatch, backend)
        r = client.get("/api/admin/dashboard")
        assert r.status_code == 200
        data = r.json()
        assert data["groups"] == {"total": 1, "active": 1, "disabled": 0}
        assert data["biz"]["checkin"] == {"done": 1, "total": 1,
                                          "unknown": False}
        assert data["biz"]["summary"] == {"done": 0, "total": 1,
                                          "unknown": False}
        assert data["biz"]["report"] == {"done": 0, "total": 1,
                                         "unknown": False}
        assert data["period_key"] == datetime.now(TZ).date().isoformat()
        assert len(data["recent"]) == 1
        assert data["recent"][0]["biz"] == "checkin"

    def test_dashboard_backend_error_is_unknown(self, tmp_db, tmp_path,
                                                monkeypatch):
        """MCP 查询失败 -> unknown=True；⛔ 不当 0 也不当成功。"""
        enable_all(tmp_db)
        tmp_db.upsert_member_task(GROUP, PHONE, "张三", None, True)
        backend = FakeBackend()
        backend.raise_on.add("checkin_query_records")
        client = logged_in_client(tmp_db, tmp_path, monkeypatch, backend)
        r = client.get("/api/admin/dashboard")
        assert r.status_code == 200
        data = r.json()
        assert data["biz"]["checkin"]["unknown"] is True
        assert data["biz"]["checkin"]["done"] == 0
        assert data["biz"]["checkin"]["total"] == 1  # 分母仍如实计

    def test_dashboard_requires_session(self, tmp_db, tmp_path, monkeypatch):
        _bootstrap(tmp_db, monkeypatch)
        client = make_client(tmp_db, tmp_path)
        assert client.get("/api/admin/dashboard").status_code == 401

    def test_dashboard_group_mode(self, tmp_db, tmp_path, monkeypatch):
        """按群模式：mode=group / biz 带 enabled / pending_people 去重计数；
        全量模式历史契约不变。"""
        enable_all(tmp_db)
        tmp_db.upsert_member_task(GROUP, PHONE, "张三", None, True)
        backend = FakeBackend()
        backend.results["checkin_query_records"] = {"records": [
            {"phone": PHONE, "name": "张三", "checkin_time":
             "2026-10-15T09:30:00+08:00", "entry_source": "user"}]}
        client = logged_in_client(tmp_db, tmp_path, monkeypatch, backend)
        r = client.get("/api/admin/dashboard?group_id=%s" % GROUP)
        assert r.status_code == 200
        data = r.json()
        assert data["ok"] is True
        assert data["mode"] == "group"
        assert data["group"]["group_id"] == GROUP
        assert data["period_key"] == datetime.now(TZ).date().isoformat()
        assert data["biz"]["checkin"]["done"] == 1
        assert data["biz"]["checkin"]["total"] == 1
        assert data["biz"]["checkin"]["enabled"] is True
        assert data["recovery"]["pending"] == 0
        assert data["reminders"]["pending"] == 0
        assert data["pending_people"] == 0
        # 未知群 -> ok False（不伪造空成功）
        r = client.get("/api/admin/dashboard?group_id=no-such-group")
        assert r.status_code == 200
        assert r.json()["ok"] is False
        # 全量模式（历史契约）仍在
        data = client.get("/api/admin/dashboard").json()
        assert data["groups"] == {"total": 1, "active": 1, "disabled": 0}
        assert data["biz"]["checkin"] == {"done": 1, "total": 1,
                                          "unknown": False}


# ---------- 成员级业务看板（2026-10-19 新增：统计中心数据源） ----------

class TestBoardApi:
    def _client(self, tmp_db, tmp_path, monkeypatch, backend):
        enable_all(tmp_db)
        tmp_db.upsert_member_task(GROUP, PHONE, "张三", "开发", True)
        return logged_in_client(tmp_db, tmp_path, monkeypatch, backend)

    def _backend(self):
        backend = FakeBackend()
        backend.results["checkin_query_records"] = {"records": [
            {"phone": PHONE, "name": "张三", "session": "am",
             "checkin_time": "2026-10-15T09:30:00+08:00",
             "status": "NORMAL", "entry_source": "user",
             "image_name": "a.jpg"}]}
        backend.results["summary_query_records"] = {"records": [
            {"phone": PHONE, "name": "张三", "content": "完成联调",
             "entry_source": "user",
             "updated_time": "2026-10-15T18:00:00+08:00"}]}
        backend.results["report_query_records"] = {"records": [
            {"phone": PHONE, "name": "张三", "entry_source": "user",
             "created_time": "2026-10-15T18:10:00+08:00",
             "items": [{"field_id": 1, "field_name": "bug修复数",
                        "unit": "个", "quantity": 3}]}]}
        return backend

    def test_board_member_matrix(self, tmp_db, tmp_path, monkeypatch):
        client = self._client(tmp_db, tmp_path, monkeypatch, self._backend())
        r = client.get("/api/admin/groups/%s/board?date=2026-10-15" % GROUP)
        assert r.status_code == 200
        data = r.json()
        assert data["ok"] is True
        assert data["date"] == "2026-10-15"
        assert [s["name"] for s in data["slots"]] == ["上午", "下午"]
        assert len(data["members"]) == 1
        member = data["members"][0]
        assert member["user_id"] == PHONE
        assert member["nickname"] == "张三"
        assert member["in_roster"] is True
        # 上午有记录 -> 完成（拍照上传）
        assert member["checkin"]["am"]["time"] == "09:30:00"
        assert member["checkin"]["am"]["status"] == "completed"
        assert member["checkin"]["am"]["method"] == "手机拍照上传"
        # 下午无记录且日期已过 -> 待补录（⛔ 不是 0 也不是成功）
        assert member["checkin"]["pm"]["status"] == "overdue"
        assert member["checkin_status"] == "overdue"
        assert member["summary"]["words"] == len("完成联调")
        assert member["summary_status"] == "completed"
        assert member["report"]["total"] == 3
        assert member["report_status"] == "completed"
        assert data["counts"]["members"] == 1
        assert data["counts"]["summary_done"] == 1
        assert data["counts"]["report_done"] == 1
        assert data["unknown"] == {"checkin": False, "summary": False,
                                   "report": False}

    def test_board_unknown_on_backend_error(self, tmp_db, tmp_path,
                                            monkeypatch):
        backend = self._backend()
        backend.raise_on.add("summary_query_records")
        client = self._client(tmp_db, tmp_path, monkeypatch, backend)
        data = client.get(
            "/api/admin/groups/%s/board?date=2026-10-15" % GROUP).json()
        assert data["unknown"]["summary"] is True
        # ⛔ 查询失败的总结不当 0 条：状态回落为待补录且不给 summary
        assert data["members"][0]["summary"] is None
        assert data["counts"]["summary_done"] == 0

    def test_board_unknown_group_404(self, tmp_db, tmp_path, monkeypatch):
        client = self._client(tmp_db, tmp_path, monkeypatch, FakeBackend())
        r = client.get("/api/admin/groups/no-such/board?date=2026-10-15")
        assert r.status_code == 404

    def test_board_bad_date_400(self, tmp_db, tmp_path, monkeypatch):
        client = self._client(tmp_db, tmp_path, monkeypatch, FakeBackend())
        r = client.get("/api/admin/groups/%s/board?date=nope" % GROUP)
        assert r.status_code == 400
        assert "日期格式不正确" in r.text

    def test_board_requires_session(self, tmp_db, tmp_path, monkeypatch):
        _bootstrap(tmp_db, monkeypatch)
        client = make_client(tmp_db, tmp_path)
        r = client.get("/api/admin/groups/%s/board" % GROUP)
        assert r.status_code == 401


# ---------- 移除群聊（2026-10-19 新增） ----------

class TestRemoveGroup:
    def _prepared(self, tmp_db, tmp_path, monkeypatch):
        enable_all(tmp_db)
        tmp_db.upsert_member_task(GROUP, PHONE, "张三", "开发", True)
        tmp_db.upsert_field(GROUP, 1, "bug修复数", None, ["开发"], True, 1)
        return logged_in_client(tmp_db, tmp_path, monkeypatch)

    def test_remove_cascades_configs(self, tmp_db, tmp_path, monkeypatch):
        client = self._prepared(tmp_db, tmp_path, monkeypatch)
        r = client.request("DELETE",
                           "/api/admin/robot-instances/%s" % GROUP,
                           headers={"X-CSRF-Token": csrf_of(client)})
        assert r.status_code == 200, r.text
        assert r.json()["ok"] is True
        # 群配置 / 名单 / 字段池全删
        assert tmp_db.get_group_config(GROUP) is None
        assert list(tmp_db.list_member_tasks(GROUP)) == []
        assert list(tmp_db.list_fields(GROUP)) == []
        # 幂等：再删一次 -> 404（⛔ 不伪装成功）
        r = client.request("DELETE",
                           "/api/admin/robot-instances/%s" % GROUP,
                           headers={"X-CSRF-Token": csrf_of(client)})
        assert r.status_code == 404

    def test_remove_requires_csrf(self, tmp_db, tmp_path, monkeypatch):
        client = self._prepared(tmp_db, tmp_path, monkeypatch)
        r = client.request("DELETE",
                           "/api/admin/robot-instances/%s" % GROUP)
        assert r.status_code == 403
        # 未删成功：配置仍在
        assert tmp_db.get_group_config(GROUP) is not None

    def test_remove_requires_session(self, tmp_db, tmp_path, monkeypatch):
        _bootstrap(tmp_db, monkeypatch)
        client = make_client(tmp_db, tmp_path)
        r = client.request("DELETE",
                           "/api/admin/robot-instances/%s" % GROUP)
        assert r.status_code == 401

    def test_removed_group_leaves_group_list(self, tmp_db, tmp_path,
                                             monkeypatch):
        client = self._prepared(tmp_db, tmp_path, monkeypatch)
        assert len(client.get("/api/admin/groups").json()["groups"]) == 1
        client.request("DELETE", "/api/admin/robot-instances/%s" % GROUP,
                       headers={"X-CSRF-Token": csrf_of(client)})
        assert client.get("/api/admin/groups").json()["groups"] == []


# ---------- 统计导出 ----------

class TestStatsExport:
    def _client(self, tmp_db, tmp_path, monkeypatch, backend):
        enable_all(tmp_db)
        tmp_db.upsert_member_task(GROUP, PHONE, "张三", None, True)
        return logged_in_client(tmp_db, tmp_path, monkeypatch, backend)

    def _records(self):
        return {"records": [
            {"phone": PHONE, "name": "张三", "session": "am",
             "checkin_time": "2026-10-15T09:30:00+08:00",
             "status": "NORMAL", "entry_source": "user",
             "image_name": "a.jpg"}]}

    def test_export_csv(self, tmp_db, tmp_path, monkeypatch):
        backend = FakeBackend()
        backend.results["checkin_query_records"] = self._records()
        client = self._client(tmp_db, tmp_path, monkeypatch, backend)
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "biz": "checkin",
                               "date": "2026-10-15", "format": "csv"})
        assert r.status_code == 200
        assert "text/csv" in r.headers["content-type"]
        body = r.content.decode("utf-8")
        # UTF-8 BOM：Excel 直接打开中文不乱码
        assert body.startswith("\ufeff")
        # 列口径与机器人侧导出一致（中文表头 + 映射值）
        assert "手机号" in body and "姓名" in body and "打卡时间" in body
        assert PHONE in body and "张三" in body
        assert "上午" in body and "按时" in body and "群员" in body
        # 中文文件名 + RFC 5987
        cd = r.headers["content-disposition"]
        assert "attachment" in cd
        assert "filename*=UTF-8''" in cd
        assert ".csv" in cd

    def test_export_xlsx(self, tmp_db, tmp_path, monkeypatch):
        backend = FakeBackend()
        backend.results["checkin_query_records"] = self._records()
        client = self._client(tmp_db, tmp_path, monkeypatch, backend)
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "biz": "checkin",
                               "date": "2026-10-15", "format": "xlsx"})
        assert r.status_code == 200
        assert r.headers["content-type"].startswith(
            "application/vnd.openxmlformats")
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(r.content))
        sheet = workbook.active
        assert sheet.title == "打卡"
        header = [cell.value for cell in sheet[1]]
        assert header[:3] == ["手机号", "姓名", "时段"]
        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        assert rows[0][0] == PHONE
        assert rows[0][1] == "张三"
        assert rows[0][2] == "上午"
        assert ".xlsx" in r.headers["content-disposition"]

    def test_export_summary_uses_group_period(self, tmp_db, tmp_path,
                                              monkeypatch):
        """总结按群周期类型：weekly 群 -> 周键（YYYY-Www）。"""
        enable_all(tmp_db)
        tmp_db.upsert_group_config(GROUP, period_type="weekly")
        backend = FakeBackend()
        backend.results["summary_query_records"] = {"records": [
            {"phone": PHONE, "name": "张三", "period_type": "weekly",
             "content": "本周正常", "entry_source": "user",
             "updated_time": "2026-10-15T18:00:00+08:00"}]}
        client = logged_in_client(tmp_db, tmp_path, monkeypatch, backend)
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "biz": "summary",
                               "date": "2026-10-15", "format": "csv"})
        assert r.status_code == 200
        body = r.content.decode("utf-8")
        assert "本周正常" in body
        # 文件名带 ISO 周键（带连字符口径，period.iso_week_key）
        from urllib.parse import unquote
        assert "2026-W42" in unquote(r.headers["content-disposition"])

    # ---------- scope=all（2026-10-19 新增：全天单表 / 三 sheet） ----------

    def test_export_all_csv_single_table(self, tmp_db, tmp_path, monkeypatch):
        backend = FakeBackend()
        backend.results["checkin_query_records"] = self._records()
        backend.results["summary_query_records"] = {"records": [
            {"phone": PHONE, "name": "张三", "content": "完成联调",
             "entry_source": "user",
             "updated_time": "2026-10-15T18:00:00+08:00"}]}
        backend.results["report_query_records"] = {"records": [
            {"phone": PHONE, "name": "张三", "entry_source": "user",
             "created_time": "2026-10-15T18:10:00+08:00",
             "items": [{"field_id": 1, "field_name": "bug修复数",
                        "unit": "个", "quantity": 3}]}]}
        client = self._client(tmp_db, tmp_path, monkeypatch, backend)
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "date": "2026-10-15",
                               "scope": "all", "format": "csv"})
        assert r.status_code == 200, r.text
        body = r.content.decode("utf-8")
        header = body.splitlines()[0]
        # 一天 × 全员 × 三业务 的单表口径（参考后台列序）
        for column in ("群聊名称", "成员姓名", "手机号码", "岗位角色",
                       "上午打卡时间", "上午打卡状态", "下午打卡时间",
                       "今日总结提交时间", "今日总结字数", "今日总结正文",
                       "今日报数提交时间", "今日报数量化指标", "报数总项数",
                       "全天综合达标判定"):
            assert column in header, column
        assert "张三" in body and "完成联调" in body and "bug修复数" in body
        assert "全天全部业务执行清单" in unquote_cd(r)

    def test_export_all_xlsx_three_sheets(self, tmp_db, tmp_path, monkeypatch):
        backend = FakeBackend()
        backend.results["checkin_query_records"] = self._records()
        client = self._client(tmp_db, tmp_path, monkeypatch, backend)
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "date": "2026-10-15",
                               "scope": "all", "format": "xlsx"})
        assert r.status_code == 200, r.text
        import openpyxl
        workbook = openpyxl.load_workbook(io.BytesIO(r.content))
        assert workbook.sheetnames == ["打卡", "总结", "报数"]

    def test_export_all_unknown_biz_503(self, tmp_db, tmp_path, monkeypatch):
        """部分业务查不到 -> 503；⛔ 不出半张空表冒充成功。"""
        backend = FakeBackend()
        backend.raise_on.add("checkin_query_records")
        client = self._client(tmp_db, tmp_path, monkeypatch, backend)
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "date": "2026-10-15",
                               "scope": "all", "format": "csv"})
        assert r.status_code == 503
        assert "稍后重试" in r.text

    def test_export_bad_scope_400(self, tmp_db, tmp_path, monkeypatch):
        client = self._client(tmp_db, tmp_path, monkeypatch, FakeBackend())
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "date": "2026-10-15",
                               "scope": "nope"})
        assert r.status_code == 400
        assert "导出范围不正确" in r.text

    def test_export_remark_goes_to_audit_only(self, tmp_db, tmp_path,
                                              monkeypatch):
        """备注只落审计存证，⛔ 不进文件内容。"""
        backend = FakeBackend()
        backend.results["checkin_query_records"] = self._records()
        client = self._client(tmp_db, tmp_path, monkeypatch, backend)
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "biz": "checkin",
                               "date": "2026-10-15", "format": "csv",
                               "remark": "例行业务归档备注"})
        assert r.status_code == 200
        assert "例行业务归档备注" not in r.content.decode("utf-8")

    def test_export_bad_biz_400(self, tmp_db, tmp_path, monkeypatch):
        client = self._client(tmp_db, tmp_path, monkeypatch, FakeBackend())
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "biz": "nope",
                               "date": "2026-10-15", "format": "csv"})
        assert r.status_code == 400
        # 通俗提示：不暴露技术名词
        assert "请选择要导出的业务" in r.text

    def test_export_bad_date_400(self, tmp_db, tmp_path, monkeypatch):
        client = self._client(tmp_db, tmp_path, monkeypatch, FakeBackend())
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "biz": "checkin",
                               "date": "not-a-date", "format": "csv"})
        assert r.status_code == 400
        assert "日期格式不正确" in r.text

    def test_export_unknown_group_404(self, tmp_db, tmp_path, monkeypatch):
        client = self._client(tmp_db, tmp_path, monkeypatch, FakeBackend())
        r = client.get("/api/admin/stats/export",
                       params={"group_id": "no-such-group",
                               "biz": "checkin", "date": "2026-10-15",
                               "format": "csv"})
        assert r.status_code == 404

    def test_export_bad_format_400(self, tmp_db, tmp_path, monkeypatch):
        client = self._client(tmp_db, tmp_path, monkeypatch, FakeBackend())
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "biz": "checkin",
                               "date": "2026-10-15", "format": "pdf"})
        assert r.status_code == 400

    def test_export_requires_session(self, tmp_db, tmp_path, monkeypatch):
        _bootstrap(tmp_db, monkeypatch)
        client = make_client(tmp_db, tmp_path)
        r = client.get("/api/admin/stats/export",
                       params={"group_id": GROUP, "biz": "checkin",
                               "date": "2026-10-15", "format": "csv"})
        assert r.status_code == 401


def unquote_cd(response):
    from urllib.parse import unquote
    return unquote(response.headers.get("content-disposition", ""))


# ---------- 文案守则（页面改由前端渲染后，改为守 API 提示与外壳） ----------

class TestCopyGuard:
    """对外文案：无 Emoji、无底层技术名词。

    页面文案已移到 admin-ui/ 源码，服务端不再产出可见文本；
    因此这里守两处仍由服务端负责的对外文案：
    1) 外壳 HTML（标题、缺失产物提示）
    2) API 错误提示（前端原样回显给管理员）
    """

    TECH_TERMS = re.compile(
        r"(sqlite|\bsql\b|json|payload|nullpointer|traceback|stack trace|"
        r"\bexception\b|\bcsrf\b|\bvalidation\b|internal server|"
        r"http\s*[45]\d\d|\bundefined\b|\bmcp\b|fail-closed|"
        r"\bwebhook\b|\bschema\b|\bnull\b)", re.IGNORECASE)

    @staticmethod
    def _assert_no_emoji(text, where):
        for ch in text:
            code = ord(ch)
            assert not (0x1F000 <= code <= 0x1FAFF
                        or 0x2600 <= code <= 0x27BF
                        or 0x2B00 <= code <= 0x2BFF
                        or code == 0xFE0F), (where, hex(code), ch)

    def test_shell_has_no_emoji(self, tmp_db, tmp_path, monkeypatch):
        client = make_client(tmp_db, tmp_path)
        self._assert_no_emoji(client.get("/login").text, "/login")

    @pytest.mark.parametrize("method,path,params,body,headers_fn", [
        ("GET", "/api/admin/stats/export",
         {"group_id": GROUP, "biz": "nope", "date": "2026-10-15"}, None, None),
        ("GET", "/api/admin/stats/export",
         {"group_id": GROUP, "date": "2026-10-15", "scope": "nope"}, None,
         None),
        ("GET", "/api/admin/groups/%s/board" % GROUP,
         {"date": "nope"}, None, None),
        ("POST", "/api/admin/robot-instances/%s/enable" % GROUP, None, None,
         "csrf"),
        ("PUT", "/api/admin/robot-instances/%s" % GROUP, None,
         {"webhook_key": "sk-x"}, "csrf"),
    ])
    def test_api_errors_are_plain(self, tmp_db, tmp_path, monkeypatch, method,
                                  path, params, body, headers_fn):
        client = logged_in_client(tmp_db, tmp_path, monkeypatch)
        headers = {"X-CSRF-Token": csrf_of(client)} if headers_fn else {}
        r = client.request(method, path, params=params, json=body,
                           headers=headers)
        assert r.status_code >= 400, r.text
        detail = r.json().get("detail") or ""
        hit = self.TECH_TERMS.search(detail)
        assert hit is None, (path, hit.group(0) if hit else "")
        self._assert_no_emoji(detail, path)
