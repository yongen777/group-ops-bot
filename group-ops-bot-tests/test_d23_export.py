# -*- coding: utf-8 -*-
"""⭐ 2026-10-20 用户指令（D23）：群内导出改「三业务三文件 xlsx」+ 全员视角。

用户当时给了五条：
1. 导出拆成**三个文件**，每个业务一个、单独命名
2. 回执只 @所有人，⛔ 不额外 @ 提问者
3. **统一 xlsx**（CSV 在量子密信电脑端打不开）
4. 每张表的字段要定义好，**看到表就一目了然今日该业务所有人的记录**
5. 「群管理员」→「后台管理员」；帮助里两个括号改「仅后台授权的管理员可用」

本文件锁的正是这五条。⛔ 与 test_stats_phase6（命令入口/权限/周期）
和 test_reminder_stats（构建器细节）不重复。
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app import commands
from app.handlers import base as base_handler
from app.handlers import bind as bind_handler
from app.handlers import realname as realname_handler
from app.handlers import stats as stats_handler
from app.web.server import create_app
from doubles.fakes import FakeBackend, RecordingSender
from helpers import GROUP, PHONE, TZ, enable_all, make_config, make_ctx

OTHER = "22644212510"
AM = datetime(2026, 10, 13, 10, 0, tzinfo=TZ)   # 落在 am 时段内（00:00-12:00）


def _setup(tmp_db, roster=(PHONE, OTHER)):
    """开全部业务 + 授权导出 + 配应交名单（enable_all 的时段 = am 00-12 / pm 12-24）。

    now=10:00 时：am **进行中** -> 未交 = 「未提交」；pm **未开始** -> 「未到时段」。
    """
    enable_all(tmp_db)
    tmp_db.upsert_whitelist(GROUP, PHONE)
    for idx, phone in enumerate(roster):
        tmp_db.upsert_member_task(GROUP, phone, "成员%d" % (idx + 1))
    return make_ctx(tmp_db, FakeBackend(),
                    make_config(Path(str(tmp_db.db_path)).parent), now=AM)


def _path(tmp_db, biz: str) -> Path:
    files = sorted((Path(str(tmp_db.db_path)).parent / "exports")
                   .glob(biz + "_*.xlsx"))
    assert files, "no %s_*.xlsx written" % biz
    return files[-1]


def _sheet(tmp_db, biz: str):
    return load_workbook(_path(tmp_db, biz)).active


class TestThreeFiles:
    """用户指令 1 + 3：三业务三文件、xlsx、业务英文 + 时间命名。"""

    def test_three_files_one_per_business(self, tmp_db):
        ctx = _setup(tmp_db)
        stats_handler.handle_command(commands.parse("导出"), ctx)
        names = sorted(p.name for p in
                       (Path(str(tmp_db.db_path)).parent / "exports").glob("*"))
        assert names == ["checkin_20261013_1000.xlsx",
                         "report_20261013_1000.xlsx",
                         "summary_20261013_1000.xlsx"]
        # ⛔ 一个 csv 都不该产生（量子密信电脑端打不开）
        assert not [n for n in names if n.endswith(".csv")]

    def test_each_file_has_one_sheet_named_by_business(self, tmp_db):
        ctx = _setup(tmp_db)
        stats_handler.handle_command(commands.parse("导出"), ctx)
        for biz, label in (("checkin", "打卡"), ("summary", "总结"),
                           ("report", "报数")):
            assert load_workbook(_path(tmp_db, biz)).sheetnames == [label]


class TestRosterView:
    """用户指令 4：全员视角 —— 没交的人也要在表里，一眼看出缺谁。"""

    def test_checkin_merges_slots_into_one_row_per_member(self, tmp_db):
        backend = FakeBackend()
        backend.results["checkin_query_records"] = {"records": [
            {"phone": PHONE, "name": "成员1", "session": "am",
             "checkin_time": "2026-10-13T09:10:00", "status": "NORMAL",
             "entry_source": "user", "image_name": "a.png"}]}
        ctx = _setup(tmp_db)
        ctx.backend = backend
        stats_handler.handle_command(commands.parse("导出"), ctx)

        ws = _sheet(tmp_db, "checkin")
        header = [c.value for c in ws[1]]
        # ⭐ 2026-10-20 用户指令：同一成员的多个时段**合并成一行**
        assert header == ["序号", "手机号", "姓名", "上午", "下午",
                          "凭证图片", "备注"]
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        assert len(rows) == 2                    # 2 人 = 2 行（⛔ 不是 2×2）
        by_phone = {row[1]: row for row in rows}   # index 1 = 手机号
        mine = by_phone[PHONE]
        assert mine[3] == "09:10:00"             # 上午已打卡（只写时刻）
        assert mine[4] == "未到打卡时间段"        # 下午还没到点
        assert mine[5] == "已提交"                # 凭证只说状态，⛔ 不显示文件名
        other = by_phone[OTHER]
        assert other[3] == "未提交"               # 成员2 上午（时段进行中）
        assert other[5] == "未提交"

    def test_missing_state_wording_and_cell_highlight(self, tmp_db):
        """⭐ 只给状态**单元格**上浅底（⛔ 不整行铺粉，否则整片粉色像坏了）。"""
        ctx = _setup(tmp_db)
        stats_handler.handle_command(commands.parse("导出"), ctx)
        ws = _sheet(tmp_db, "checkin")
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                if cell.value in ("未提交", "未到打卡时间段"):
                    assert str(cell.fill.fgColor.rgb).endswith("FDE8E8"), \
                        cell.coordinate
                else:
                    assert str(cell.fill.fgColor.rgb).endswith("00000000"), \
                        (cell.coordinate, cell.value)

    def test_summary_and_report_list_absent_members(self, tmp_db):
        ctx = _setup(tmp_db)      # 无任何记录 -> 全员未提交
        stats_handler.handle_command(commands.parse("导出"), ctx)
        for biz in ("summary", "report"):
            rows = list(_sheet(tmp_db, biz).iter_rows(min_row=2,
                                                      values_only=True))
            assert len(rows) == 2, biz
            assert {r[1] for r in rows} == {"未提交"}, biz
            assert {r[2] for r in rows} == {PHONE, OTHER}, biz

    def test_record_outside_roster_is_still_listed(self, tmp_db):
        """不在应交名单但有记录的人 ⛔ 不能被静默丢掉（数据不能消失）。"""
        backend = FakeBackend()
        backend.results["summary_query_records"] = {"records": [
            {"phone": "22600000009", "name": "外部人", "period_type": "daily",
             "content": "临时支援", "updated_time": "20:00:00",
             "entry_source": "user"}]}
        ctx = _setup(tmp_db)
        ctx.backend = backend
        stats_handler.handle_command(commands.parse("导出"), ctx)

        rows = list(_sheet(tmp_db, "summary").iter_rows(min_row=2,
                                                        values_only=True))
        assert len(rows) == 3
        extra = [r for r in rows if r[2] == "22600000009"]
        assert extra and extra[0][8] == "不在应交名单"

    def test_report_quantity_is_numeric(self, tmp_db):
        """数量必须是数字（⛔ 文本会让 Excel 求和变 0）。"""
        backend = FakeBackend()
        backend.results["report_query_records"] = {"records": [
            {"phone": PHONE, "name": "成员1", "created_time": "20:00:00",
             "entry_source": "user",
             "items": [{"field_id": 1, "field_name": "bug修复数",
                        "quantity": 3.0, "unit": None}]}]}
        ctx = _setup(tmp_db)
        ctx.backend = backend
        stats_handler.handle_command(commands.parse("导出"), ctx)
        ws = _sheet(tmp_db, "report")
        header = [c.value for c in ws[1]]
        cell = ws["F2"].value          # 序号/状态/手机/姓名/时间 + 字段1
        assert header[5] == "bug修复数"
        assert isinstance(cell, (int, float)) and cell == 3


class TestColumnIntegrity:
    """⚠️ 线上实测踩过：未提交行多写一个空串 -> 表头右端冒出一个空列。

    Excel 里那一列既没标题也没内容，筛选/求和全歪。这里锁「表头无空洞、
    最大列数 = 表头列数」—— 行比表头宽时 openpyxl 会给表头补 None。
    """

    def test_header_has_no_hole_and_matches_max_column(self, tmp_db):
        backend = FakeBackend()
        backend.results["report_query_records"] = {"records": [
            {"phone": PHONE, "name": "成员1", "created_time": "20:00:00",
             "entry_source": "user",
             "items": [{"field_id": 1, "field_name": "bug修复数",
                        "quantity": 3.0, "unit": None},
                       {"field_id": 2, "field_name": "代码行数",
                        "quantity": 500, "unit": "行"}]}]}
        ctx = _setup(tmp_db)
        ctx.backend = backend
        stats_handler.handle_command(commands.parse("导出"), ctx)
        for biz in ("checkin", "summary", "report"):
            ws = _sheet(tmp_db, biz)
            header = [c.value for c in ws[1]]
            assert None not in header, (biz, header)
            assert ws.max_column == len(header), (biz, ws.max_column,
                                                  len(header))


class TestBroadcastReceipt:
    """用户指令 2：导出成功回执只 @所有人，⛔ 不额外 @ 提问者。"""

    def _app(self, tmp_db, file_ok: bool):
        sender = RecordingSender()
        sender.file_ok = file_ok
        app = create_app(
            state=tmp_db, backend=FakeBackend(), sender=sender,
            config=make_config(Path(str(tmp_db.db_path)).parent),
            now_fn=lambda: AM)
        return TestClient(app), sender

    def _post_export(self, client):
        return client.post("/api/callback", json={
            "group_id": GROUP, "phone": PHONE, "text": "导出",
            "reply_url": "http://reply.test/hook", "name": "测试用户"})

    def test_success_receipt_mentions_nobody(self, tmp_db):
        _setup(tmp_db)
        client, sender = self._app(tmp_db, True)
        resp = self._post_export(client)
        assert resp.json()["replied"] is True
        assert len(sender.sent) == 1
        assert sender.sent[0]["text"] == stats_handler.EXPORT_SENT_TEXT
        # ⛔ 不再带 note 提问者的手机号（文案里已是 @所有人）
        assert sender.sent[0]["mention_phones"] == []

    def test_failure_receipt_still_mentions_asker(self, tmp_db):
        """失败要 @ 提问者 —— 他得知道文件没发出去。"""
        _setup(tmp_db)
        client, sender = self._app(tmp_db, False)
        self._post_export(client)
        assert sender.sent[0]["text"] == stats_handler.EXPORT_SEND_FAILED_TEXT
        assert sender.sent[0]["mention_phones"] == [PHONE]

    def test_broadcast_predicate_is_text_exact(self):
        assert stats_handler.is_broadcast_receipt(
            stats_handler.EXPORT_SENT_TEXT)
        # 失败 / 部分成功⛔ 不算群发回执（必须 @ 提问者）
        assert not stats_handler.is_broadcast_receipt(
            stats_handler.EXPORT_SEND_FAILED_TEXT)
        assert not stats_handler.is_broadcast_receipt("随便一句话")


class TestAdminWording:
    """用户指令 5：群管理员 ≠ 后台管理员，回复文本一律说「后台管理员」。"""

    def test_help_text_uses_backend_admin_qualifier(self):
        assert "（仅后台授权的管理员可用）" in commands.HELP_TEXT
        assert commands.HELP_TEXT.count("（仅后台授权的管理员可用）") == 2
        assert "（仅管理员可用）" not in commands.HELP_TEXT
        # 绑定成功后的命令清单与帮助文案同源（⛔ 不能一处改一处漏）
        assert "（仅后台授权的管理员可用）" in \
            bind_handler.BOUND_COMMANDS_GUIDE

    def test_group_facing_texts_never_say_group_admin(self):
        texts = {
            "PAUSED_MEMBER_TEXT": base_handler.PAUSED_MEMBER_TEXT,
            "UNBOUND_TEXT": bind_handler.UNBOUND_TEXT,
            "DISABLED_GROUP_TEXT": bind_handler.DISABLED_GROUP_TEXT,
            "ALREADY_REGISTERED_TEXT":
                realname_handler.ALREADY_REGISTERED_TEXT,
            "HELP_TEXT": commands.HELP_TEXT,
            "BOUND_COMMANDS_GUIDE": bind_handler.BOUND_COMMANDS_GUIDE,
        }
        for name, text in texts.items():
            assert "群管理员" not in text, name
        assert "后台管理员" in base_handler.PAUSED_MEMBER_TEXT
        # ⭐ 2026-10-20（D28）：绑定对全体群成员开放 —— 未绑定群引导语不再让
        # 用户去找「后台管理员」，改成明确「本群任一成员均可操作」。
        assert "任一成员" in bind_handler.UNBOUND_TEXT
        assert "后台管理员" in bind_handler.DISABLED_GROUP_TEXT
